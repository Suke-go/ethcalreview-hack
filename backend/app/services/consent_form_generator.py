"""
同意書生成サービス（研究計画概要付き）

フォームデータから研究計画概要を含む同意書をdocx形式で生成します。
裏面「研究の概要について」の各項目は、フォームが空の場合 LLM で自動補完します。
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
from typing import Dict, Any, Optional
from datetime import datetime

from app.services.llm_client import LLMClient
from app.services.academic_writing_generator import generate_research_plan_sections
from app.logger import get_logger

logger = get_logger(__name__)


# LLM が使えない/失敗した場合のフォールバック文言
_FALLBACK_PLACEHOLDERS: Dict[str, str] = {
    "participant_criteria": (
        "成人（18歳以上）で、本研究の趣旨について十分な説明を受け、参加に同意の得られた健常者を対象とする。"
        "重篤な疾患を有する者、および研究内容の理解が困難な者は除外される。"
        "研究への参加は任意であり、不参加によって不利益が生じることはない。"
    ),
    "research_purpose": (
        "本研究は、研究課題に関連する現象について実験的に検証し、その特性を明らかにすることを目的とする。"
    ),
    "research_significance": (
        "本研究で得られる知見は、関連する学術分野の理解に資するとともに、"
        "今後の応用研究の基礎として貢献することが期待される。"
    ),
    "research_method": (
        "実験室環境において、研究対象者に所定の課題を遂行してもらい、必要な計測および記録を行う。"
        "実験手続きは研究計画に従って統制される。"
    ),
    "risks": (
        "本研究において、日常生活で生じる程度を超えるリスクは想定されない。"
        "万一、研究対象者が体調不良を訴えた場合は直ちに実験を中止し、適切に対応する。"
    ),
}


async def generate_consent_form_with_plan(
    form_data: Dict[str, Any],
    output_dir: Path,
    lab_defaults: Dict[str, Any],
    llm_client: Optional[LLMClient] = None,
) -> Path:
    """
    研究計画概要を含む同意書を生成

    Args:
        form_data: フォームデータ
        output_dir: 出力ディレクトリ
        lab_defaults: 研究室デフォルト設定
        llm_client: LLM クライアント（None の場合はフォールバック文言を使用）

    Returns:
        生成されたファイルのパス
    """
    # 裏面「研究の概要について」の各項目を、未入力なら自動補完
    form_data = await _ensure_back_side_sections(form_data, lab_defaults, llm_client)

    doc = Document()
    
    # スタイル設定（11pt統一）
    style = doc.styles['Normal']
    style.font.name = 'Yu Gothic'
    style.font.size = Pt(11)
    
    # ========================================
    # 表面（基本情報）
    # ========================================
    
    # タイトル
    title = doc.add_paragraph()
    title_run = title.add_run('同意書')
    title_run.bold = True
    title_run.font.size = Pt(16)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # 研究課題名
    doc.add_paragraph(f"研究課題名: {form_data.get('research_title', '')}")
    
    # 研究責任者
    pi_name = lab_defaults.get('lab_info', {}).get('pi_name', '')
    pi_affiliation = lab_defaults.get('lab_info', {}).get('pi_affiliation', '')
    pi_tel = lab_defaults.get('lab_info', {}).get('pi_phone', '')
    
    doc.add_paragraph(f"研究責任者: {pi_affiliation} {pi_name}")
    doc.add_paragraph(f"連絡先: {pi_tel}")
    
    doc.add_paragraph()
    
    # 同意文
    consent_text = doc.add_paragraph()
    consent_text.add_run(
        f"私は、上記の研究について十分な説明を受け、研究の内容を理解しました。\n"
        f"自分の自由意思に基づき、本研究に参加することに同意します。"
    )
    
    doc.add_paragraph()
    
    # 署名欄
    doc.add_paragraph(f"同意日: {datetime.now().strftime('%Y年%m月%d日')}")
    doc.add_paragraph("研究対象者氏名: ___________________________")
    doc.add_paragraph("署名: ___________________________")
    
    # ページブレーク
    doc.add_page_break()
    
    # ========================================
    # 裏面（研究の概要）
    # ========================================
    
    overview_title = doc.add_paragraph()
    overview_run = overview_title.add_run('研究の概要について')
    overview_run.bold = True
    overview_run.font.size = Pt(14)
    overview_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # 研究対象者条件
    _add_section(doc, "研究対象者条件", form_data.get('participant_criteria', ''))
    
    # 目的
    _add_section(doc, "目的", form_data.get('research_purpose', ''))
    
    # 意義
    _add_section(doc, "意義", form_data.get('research_significance', ''))
    
    # 方法
    _add_section(doc, "方法", form_data.get('research_method', ''))
    
    # 実験内容（手順）
    if 'experiment_procedures' in form_data:
        procedures_heading = doc.add_paragraph()
        procedures_heading.add_run('[実験内容]').bold = True
        
        procedures = form_data['experiment_procedures']
        if isinstance(procedures, list):
            for i, proc in enumerate(procedures, 1):
                proc_text = proc if isinstance(proc, str) else proc.get('description', '')
                doc.add_paragraph(f"{i}. {proc_text}", style='List Number')
        else:
            doc.add_paragraph(str(procedures))
        
        doc.add_paragraph()
    
    # 所要時間
    duration = form_data.get('duration_minutes', '')
    if duration:
        _add_section(doc, "所要時間", f"約{duration}分")
    
    # 考えられるリスク
    _add_section(doc, "考えられるリスク", form_data.get('risks', ''))
    
    # 謝礼
    reward_amount = form_data.get('reward_amount', '')
    reward_type = form_data.get('reward_type', 'Amazonギフトカード（Eメールタイプ）')
    if reward_amount:
        _add_section(doc, "謝礼", 
                    f"実験終了時に謝金を{reward_type}にて{reward_amount}円お支払いいたします。")
    
    # 研究対象者の必要性、リスクと安全性
    _add_section(doc, "研究対象者の必要性，研究への参加におけるリスクと安全性，危険回避の方法について",
                form_data.get('safety_measures', ''))
    
    # 個人情報の保護
    privacy_heading = doc.add_paragraph()
    privacy_heading.add_run('個人情報の保護について').bold = True
    
    # データの匿名化
    anonymization = form_data.get('anonymization_method',
        '連結可能匿名化によるデータの匿名化を行います。研究対象者の氏名等の個人情報は研究データとは分離して管理します。')
    doc.add_paragraph(f"(1) データの匿名化について\n{anonymization}")
    
    # データの管理方法
    data_management = lab_defaults.get('data_management_defaults', {})
    management_location = data_management.get('management_location', '')
    management_method = data_management.get('management_method', '')
    disposal_method = data_management.get('disposal_method', '')
    
    doc.add_paragraph(
        f"(2) データの管理方法\n"
        f"研究データは{management_location}で保管します。{management_method}\n"
        f"{disposal_method}"
    )
    
    # 研究参加の任意性
    consent_withdrawal = form_data.get('consent_withdrawal_deadline', '同意書署名の日から90日後')
    doc.add_paragraph(
        f"(3) 研究参加の任意性\n"
        f"研究への参加は任意であり、参加しないことで不利益が生じることはありません。"
        f"実験の途中であっても不利益なく参加を取りやめることができます。"
        f"また、{consent_withdrawal}までであればデータ提供の同意撤回が可能です。"
    )
    
    # 保存
    output_path = output_dir / "03_同意書（サンプル）.docx"
    doc.save(output_path)
    
    return output_path


def _add_section(doc: Document, title: str, content: str):
    """セクションを追加するヘルパー関数"""
    if not content:
        return

    heading = doc.add_paragraph()
    heading.add_run(f'[{title}]').bold = True

    # 改行を保持してパラグラフ追加
    for line in content.split('\n'):
        if line.strip():
            doc.add_paragraph(line.strip())

    doc.add_paragraph()


# ============================================================
# 裏面「研究の概要について」の自動補完ロジック
# ============================================================

# LLM で生成する記述系フィールド（一括生成）
_AI_TEXT_FIELDS = (
    "participant_criteria",
    "research_purpose",
    "research_significance",
    "research_method",
    "risks",
)


def _is_blank(value: Any) -> bool:
    """フォーム値が空(未入力)かどうかを判定"""
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, (list, dict, tuple, set)):
        return len(value) == 0
    return False


def _estimate_reward_amount(duration_minutes: Any) -> int:
    """所要時間から大まかな謝礼額を見積もる（時給1000円換算 / 30分刻み 500円）"""
    try:
        minutes = int(duration_minutes)
    except (TypeError, ValueError):
        minutes = 60
    return max(500, int(round(minutes / 30) * 500))


def _brief_description_from_form(form: Dict[str, Any]) -> str:
    """LLM 生成の入力となる『簡単な説明』をフォームから組み立てる"""
    candidates = [
        form.get("brief_description"),
        form.get("research_overview"),
        form.get("background"),
        form.get("purpose"),
        form.get("experiment_objective"),
    ]
    parts = [str(c).strip() for c in candidates if c and str(c).strip()]
    return "\n".join(parts) if parts else (form.get("research_title") or "")


async def _ensure_back_side_sections(
    form_data: Dict[str, Any],
    lab_defaults: Dict[str, Any],
    llm_client: Optional[LLMClient],
) -> Dict[str, Any]:
    """
    同意書裏面「研究の概要について」の各項目を、未入力ならLLMで自動補完する。
    すでに値があるフィールドは上書きしない。
    """
    form = dict(form_data)

    # ---- 1) 記述系 5 項目: 空のものをまとめて LLM 生成 ----
    missing_fields = [f for f in _AI_TEXT_FIELDS if _is_blank(form.get(f))]

    if missing_fields:
        ai_generated: Dict[str, str] = {}
        if llm_client is not None:
            try:
                logger.info(f"裏面の未入力セクションを LLM で補完: {missing_fields}")
                ai_generated = await generate_research_plan_sections(
                    research_title=form.get("research_title") or "本研究",
                    brief_description=_brief_description_from_form(form),
                    research_field=form.get("research_field") or "ヒューマンインタフェース",
                    llm_client=llm_client,
                )
            except Exception as e:
                logger.warning(
                    f"LLM による裏面セクション自動生成に失敗、フォールバック文言を使用: "
                    f"{type(e).__name__}: {e}"
                )
                ai_generated = {}
        else:
            logger.info("LLM クライアント未設定。裏面セクションはフォールバック文言で補完")

        for field in missing_fields:
            text = (ai_generated.get(field) or "").strip()
            if not text:
                text = _FALLBACK_PLACEHOLDERS[field]
            form[field] = text

    # ---- 2) 所要時間: 未入力なら 60 分 ----
    if _is_blank(form.get("duration_minutes")):
        form["duration_minutes"] = 60

    # ---- 3) 謝礼: 未入力なら lab_defaults / 既定値で補完 ----
    reward_defaults = (lab_defaults or {}).get("reward_defaults", {}) or {}
    if _is_blank(form.get("reward_amount")):
        form["reward_amount"] = (
            reward_defaults.get("default_amount")
            or _estimate_reward_amount(form.get("duration_minutes"))
        )
    if _is_blank(form.get("reward_type")):
        form["reward_type"] = (
            reward_defaults.get("default_type")
            or "Amazonギフトカード（Eメールタイプ）"
        )

    return form
