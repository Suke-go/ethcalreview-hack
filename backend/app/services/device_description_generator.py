"""
新規開発デバイス説明書 生成サービス

新規開発デバイスを使用する研究の場合に、
デバイスの詳細説明と安全性資料を生成します。
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
from typing import Dict, Any, List
from datetime import datetime

from app.services.llm_client import LLMClient
from app.logger import get_logger

logger = get_logger(__name__)


DEVICE_DESCRIPTION_PROMPT = """
あなたは研究倫理審査の専門家です。
以下の研究で使用する新規開発デバイスについて、倫理審査申請書に添付する説明資料を作成してください。

【研究情報】
研究タイトル: {research_title}
研究概要: {research_purpose}
研究方法: {research_method}
使用機器: {devices}
想定されるリスク: {risks}
リスク対策: {risk_countermeasures}

【執筆ルール】
1. である調で統一する
2. 専門用語は必要最小限とし、使用する場合は平易な説明を付ける
3. 安全性を裏付ける根拠を具体的に記述する
4. すべて断定形で記述する（「〜と考えられる」ではなく「〜である」）

以下の5セクションを出力してください。各セクションは【見出し】で始めてください。

【デバイスを使用する目的】
- このデバイスを研究で使用する理由と目的
- 研究課題との関連性
- 200-300字で記述

【開発した新規機能】
- デバイスの新規機能の説明
- 従来技術との違い
- 200-300字で記述

【システムの構成】
- デバイスを構成する要素
- 各要素の役割と接続関係
- 300-400字で記述

【機器が発生する意図しない刺激と回避策】
- 想定される意図しない刺激（リスク1、リスク2...の形式で列挙）
- 各リスクに対する具体的な回避策（インデント付きで記述）
- 400-500字で記述

【被験者が受ける最大の刺激が安全であることを裏付ける資料】
- 最大刺激条件の定義
- 安全性の根拠（科学的根拠、先行研究、設計上の安全対策など）
- フェイルセーフ機構の説明
- 試行時間や休憩の設計による安全確保
- 500-600字で記述
"""


class DeviceDescriptionGenerator:
    """新規開発デバイス説明書 生成器"""
    
    def __init__(self, llm_client: LLMClient, lab_defaults: Dict[str, Any]):
        self.llm = llm_client
        self.defaults = lab_defaults
    
    async def generate(
        self,
        form_data: Dict[str, Any],
        output_dir: Path
    ) -> Path:
        """
        デバイス説明書を生成
        """
        logger.info("=" * 60)
        logger.info("新規開発デバイス説明書生成 開始")
        
        # コンテキスト準備
        context = self._build_context(form_data)
        
        # LLMで説明文を生成
        description_text = await self._generate_description(context)
        
        # DOCXファイル生成
        output_path = self._generate_docx(description_text, context, output_dir)
        
        logger.info(f"新規開発デバイス説明書生成 完了: {output_path.name}")
        logger.info("=" * 60)
        
        return output_path
    
    def _build_context(self, form_data: Dict[str, Any]) -> Dict[str, Any]:
        """フォームデータからコンテキストを構築"""
        devices = form_data.get("devices", [])
        risks = form_data.get("risks", [])
        risk_countermeasures = form_data.get("riskCountermeasures", form_data.get("risk_countermeasures", []))
        
        return {
            "research_title": form_data.get("title", form_data.get("research_title", "")),
            "research_purpose": form_data.get("purpose", form_data.get("research_purpose", "")),
            "research_method": form_data.get("methodology", form_data.get("research_method", "")),
            "devices": ", ".join(devices) if devices else "特になし",
            "risks": "\n".join(f"- {r}" for r in risks) if risks else "特になし",
            "risk_countermeasures": "\n".join(f"- {c}" for c in risk_countermeasures) if risk_countermeasures else "",
        }
    
    async def _generate_description(self, context: Dict[str, Any]) -> str:
        """LLMでデバイス説明を生成"""
        prompt = DEVICE_DESCRIPTION_PROMPT.format(**context)
        
        try:
            response = await self.llm.generate_content_async(
                prompt=prompt,
                system_instruction=(
                    "あなたは研究倫理審査申請書を作成する経験豊富な研究者です。"
                    "新規開発デバイスの説明と安全性の根拠を、審査委員に分かりやすく説明してください。"
                    "すべて断定形で書き、曖昧な表現は避けてください。"
                    "である調で統一してください。"
                )
            )
            return response.strip()
        except Exception as e:
            logger.error(f"デバイス説明生成エラー: {e}")
            return self._generate_fallback(context)
    
    def _generate_fallback(self, context: Dict[str, Any]) -> str:
        """LLM失敗時のフォールバック"""
        return f"""
【デバイスを使用する目的】
本研究では、{context['devices']}を使用する。
{context['research_purpose']}を達成するため、本デバイスを用いて実験を行う。

【開発した新規機能】
（研究者が記入してください）

【システムの構成】
（研究者が記入してください）

【機器が発生する意図しない刺激と回避策】
想定されるリスク：
{context['risks']}

回避策：
{context['risk_countermeasures']}

【被験者が受ける最大の刺激が安全であることを裏付ける資料】
（研究者が記入してください）
"""
    
    def _generate_docx(
        self,
        description_text: str,
        context: Dict[str, Any],
        output_dir: Path
    ) -> Path:
        """説明書DOCXを生成"""
        doc = Document()
        
        # スタイル設定
        style = doc.styles['Normal']
        style.font.name = 'Yu Gothic'
        style.font.size = Pt(11)
        
        # タイトル
        title = doc.add_paragraph()
        title_run = title.add_run("新規開発デバイス説明書")
        title_run.bold = True
        title_run.font.size = Pt(14)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        # 研究課題名
        subtitle = doc.add_paragraph()
        subtitle.add_run(f"研究課題名: {context['research_title']}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        doc.add_paragraph()
        
        # 説明文をパースして追加
        lines = description_text.split("\n")
        current_section = None
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # 見出しの検出
            if line.startswith("【") and line.endswith("】"):
                doc.add_paragraph()
                heading = doc.add_paragraph()
                heading_run = heading.add_run(line)
                heading_run.bold = True
                current_section = line
            # 回避策（インデント付き）
            elif line.startswith("回避策") or line.startswith("対策"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                run = p.add_run(line)
            # 箇条書き項目
            elif line.startswith("-") or line.startswith("・"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run(line)
            # 通常の段落
            else:
                doc.add_paragraph(line)
        
        # 注記
        doc.add_paragraph()
        doc.add_paragraph()
        note = doc.add_paragraph()
        note.add_run("※本資料は研究倫理審査申請書の添付資料として作成された。")
        
        # 保存
        output_path = output_dir / "新規開発デバイス説明書.docx"
        doc.save(output_path)
        
        return output_path


async def generate_device_description(
    form_data: Dict[str, Any],
    output_dir: Path,
    llm_client: LLMClient,
    lab_defaults: Dict[str, Any]
) -> Path:
    """
    新規開発デバイス説明書生成エントリーポイント
    """
    generator = DeviceDescriptionGenerator(llm_client, lab_defaults)
    return await generator.generate(form_data, output_dir)
