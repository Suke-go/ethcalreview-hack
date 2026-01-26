"""
ドキュメント生成サービス
"""
from pathlib import Path
from typing import Dict, Any, List, Optional
from datetime import datetime
import zipfile
import io
from docxtpl import DocxTemplate
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.config import UserSettings
from app.logger import get_logger

logger = get_logger(__name__)

class DocumentGenerator:
    """docxtplを使用してドキュメントを生成"""
    
    TEMPLATE_MAP = {
        "application_form": "01-1_申請書_template.docx",
        "consent_form": "03_同意書_template.docx",
        "consent_withdrawal": "04_同意撤回書_template.docx",
        "implementation_plan": "01-2_実施計画書_template.docx",
        "recruitment_notice": "募集案内文_template.docx",
    }
    
    OUTPUT_NAMES = {
        "application_form": "01-1_研究倫理審査申請書.docx",
        "consent_form": "03_同意書.docx",
        "consent_withdrawal": "04_同意撤回書.docx",
        "implementation_plan": "01-2_実施計画書.docx",
        "recruitment_notice": "添付資料B_実験参加者募集案内文.docx",
        "questionnaire": "添付資料C_アンケート用紙.docx",
        "experiment_script": "添付資料D_実験説明台本.docx",
        "participant_list": "参加者リスト.xlsx",
    }
    
    def __init__(
        self, 
        templates_dir: Path, 
        output_dir: Path,
        settings: UserSettings
    ):
        self.templates_dir = templates_dir
        self.output_dir = output_dir
        self.settings = settings
        self.output_dir.mkdir(parents=True, exist_ok=True)
    
    def _build_context(self, form_data: Dict[str, Any]) -> Dict[str, Any]:
        """フォームデータと設定を統合してテンプレートコンテキストを構築"""
        # 設定からのデフォルト値
        context = {
            # 研究責任者
            "pi_name": self.settings.principal_investigator.name,
            "pi_affiliation": self.settings.principal_investigator.affiliation,
            "pi_position": self.settings.principal_investigator.position,
            "pi_email": self.settings.principal_investigator.email,
            "pi_tel": self.settings.principal_investigator.phone,
            
            # 実施担当者
            "conductor_tel": self.settings.experiment_conductor.phone,
            
            # 倫理委員会
            "ethics_committee": self.settings.ethics_committee.name,
            "ethics_office": self.settings.ethics_committee.office,
            "ethics_tel": self.settings.ethics_committee.phone,
            
            # 研究室
            "lab_name": self.settings.laboratory.name,
            "lab_building": self.settings.laboratory.building,
            "lab_room": self.settings.laboratory.room,
            
            # 予算
            "reward_amount": self.settings.budget.reward_per_person,
            "reward_type": self.settings.budget.reward_type,
            "budget_source": self.settings.budget.source,
            
            # 保険
            "insurance_type": self.settings.insurance.type,
            
            # 日付
            "current_date": datetime.now().strftime("%Y年%m月%d日"),
            "current_year": datetime.now().year,
        }
        
        # フォームデータで上書き
        context.update(form_data)
        
        # チェックボックス用のヘルパー関数を追加
        context["checkbox"] = lambda v: "■" if v else "□"
        
        return context
    
    def generate(self, doc_type: str, form_data: Dict[str, Any]) -> Path:
        """ドキュメントを生成"""
        logger.debug(f"generate() 開始: {doc_type}")
        
        if doc_type not in self.TEMPLATE_MAP:
            logger.error(f"不明なドキュメントタイプ: {doc_type}")
            raise ValueError(f"Unknown document type: {doc_type}")
        
        template_path = self.templates_dir / self.TEMPLATE_MAP[doc_type]
        
        if not template_path.exists():
            logger.error(f"テンプレートが見つかりません: {template_path}")
            raise FileNotFoundError(f"Template not found: {template_path}")
        
        # テンプレート読み込み
        logger.debug(f"  テンプレート読み込み: {template_path.name}")
        doc = DocxTemplate(template_path)
        
        # コンテキスト構築
        context = self._build_context(form_data)
        logger.debug(f"  コンテキスト: {len(context)} フィールド")
        
        # Jinja2環境を作成してフィルターを登録
        from jinja2 import Environment
        jinja_env = Environment()
        jinja_env.filters["checkbox"] = lambda v: "■" if v else "□"
        
        # レンダリング
        logger.debug("  レンダリング実行中...")
        doc.render(context, jinja_env=jinja_env)
        
        # 保存
        output_path = self.output_dir / self.OUTPUT_NAMES[doc_type]
        doc.save(output_path)
        logger.debug(f"  保存完了: {output_path.name}")
        
        return output_path
    
    def generate_questionnaire(
        self, 
        title: str,
        questions: List[Dict[str, Any]],
        form_data: Dict[str, Any]
    ) -> Path:
        """アンケート用紙を新規生成（python-docx）"""
        doc = Document()
        
        # スタイル設定
        style = doc.styles['Normal']
        style.font.name = 'Yu Gothic'
        style.font.size = Pt(10.5)
        
        # タイトル
        title_para = doc.add_paragraph()
        title_run = title_para.add_run("アンケート用紙")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 研究タイトル
        research_title = doc.add_paragraph()
        research_title.add_run(f"研究課題名: {title}")
        research_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()  # 空行
        
        # 説明文
        intro = doc.add_paragraph()
        intro.add_run("以下の質問にお答えください。回答に正解・不正解はありません。")
        
        doc.add_paragraph()  # 空行
        
        # 質問項目
        for i, q in enumerate(questions, 1):
            q_para = doc.add_paragraph()
            q_para.add_run(f"Q{i}. {q.get('text', '')}").bold = True
            
            q_type = q.get('type', 'text')
            
            if q_type == 'likert':
                # リカート尺度
                scale = q.get('scale', 5)
                labels = q.get('labels', ['全くそう思わない', '', '', '', '非常にそう思う'])
                
                table = doc.add_table(rows=2, cols=scale)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # ラベル行
                for j, label in enumerate(labels):
                    cell = table.rows[0].cells[j]
                    cell.text = label
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # 選択肢行
                for j in range(scale):
                    cell = table.rows[1].cells[j]
                    cell.text = f"□ {j + 1}"
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
            elif q_type == 'choice':
                # 選択式
                options = q.get('options', [])
                for opt in options:
                    opt_para = doc.add_paragraph()
                    opt_para.add_run(f"  □ {opt}")
                    
            elif q_type == 'text':
                # 自由記述
                doc.add_paragraph("_" * 60)
                doc.add_paragraph("_" * 60)
            
            doc.add_paragraph()  # 空行
        
        # 終了メッセージ
        end_para = doc.add_paragraph()
        end_para.add_run("ご協力ありがとうございました。").bold = True
        end_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 保存
        output_path = self.output_dir / self.OUTPUT_NAMES["questionnaire"]
        doc.save(output_path)
        
        return output_path
    
    def generate_experiment_script(
        self, 
        title: str,
        procedures: List[Dict[str, Any]],
        form_data: Dict[str, Any]
    ) -> Path:
        """実験説明台本を新規生成（python-docx）"""
        doc = Document()
        
        # スタイル設定
        style = doc.styles['Normal']
        style.font.name = 'Yu Gothic'
        style.font.size = Pt(10.5)
        
        # タイトル
        title_para = doc.add_paragraph()
        title_run = title_para.add_run("実験説明台本")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 研究タイトル
        research_title = doc.add_paragraph()
        research_title.add_run(f"研究課題名: {title}")
        research_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()  # 空行
        
        # 注意書き
        note = doc.add_paragraph()
        note.add_run("※ 以下は実験実施者が参加者に説明する際の台本です。").italic = True
        
        doc.add_paragraph()  # 空行
        
        # 挨拶
        greeting = doc.add_paragraph()
        greeting.add_run("【挨拶】").bold = True
        doc.add_paragraph(
            f"「本日は{form_data.get('lab_name', '本研究室')}の実験にご参加いただき、"
            "誠にありがとうございます。私は本実験の担当者の○○と申します。"
            "これから実験の内容について説明させていただきます。」"
        )
        
        doc.add_paragraph()
        
        # 研究目的
        purpose_title = doc.add_paragraph()
        purpose_title.add_run("【研究目的の説明】").bold = True
        doc.add_paragraph(
            f"「本研究は、{form_data.get('purpose', '～を目的として')}行います。"
            f"{form_data.get('methodology', '')}」"
        )
        
        doc.add_paragraph()
        
        # 実験手順
        procedure_title = doc.add_paragraph()
        procedure_title.add_run("【実験手順の説明】").bold = True
        
        for i, proc in enumerate(procedures, 1):
            proc_para = doc.add_paragraph()
            proc_para.add_run(f"手順{i}: ").bold = True
            proc_para.add_run(proc.get('description', ''))
            
            if proc.get('duration'):
                proc_para.add_run(f" （所要時間: 約{proc['duration']}分）")
        
        doc.add_paragraph()
        
        # 注意事項
        caution_title = doc.add_paragraph()
        caution_title.add_run("【注意事項の説明】").bold = True
        
        cautions = [
            "実験中、体調が悪くなったり、続けることが難しいと感じた場合は、遠慮なくお申し出ください。",
            "実験はいつでも中止することができ、中止しても不利益を受けることはありません。",
            "実験中に撮影・収集したデータは、研究目的以外には使用いたしません。",
        ]
        
        for caution in cautions:
            doc.add_paragraph(f"・{caution}")
        
        doc.add_paragraph()
        
        # 同意確認
        consent_title = doc.add_paragraph()
        consent_title.add_run("【同意の確認】").bold = True
        doc.add_paragraph(
            "「以上が本実験の説明となります。ご不明な点はございますか？"
            "よろしければ、同意書にご署名をお願いいたします。」"
        )
        
        doc.add_paragraph()
        
        # 終了時
        end_title = doc.add_paragraph()
        end_title.add_run("【実験終了時】").bold = True
        reward = form_data.get('reward_amount', 1000)
        doc.add_paragraph(
            f"「以上で実験は終了です。ご協力ありがとうございました。"
            f"謝礼として{reward:,}円をお渡しいたします。"
            "受領書にご署名をお願いいたします。」"
        )
        
        # 保存
        output_path = self.output_dir / self.OUTPUT_NAMES["experiment_script"]
        doc.save(output_path)
        
        return output_path
    
    def generate_all(self, form_data: Dict[str, Any]) -> Dict[str, Optional[Path]]:
        """すべてのドキュメントを生成"""
        results: Dict[str, Optional[Path]] = {}
        
        # テンプレートベースの書類
        for doc_type in self.TEMPLATE_MAP:
            try:
                path = self.generate(doc_type, form_data)
                results[doc_type] = path
            except Exception as e:
                results[doc_type] = None
                print(f"Error generating {doc_type}: {e}")
        
        # アンケート用紙（質問がある場合）
        if form_data.get('questionnaire_items'):
            try:
                path = self.generate_questionnaire(
                    title=form_data.get('research_title', ''),
                    questions=form_data['questionnaire_items'],
                    form_data=form_data
                )
                results['questionnaire'] = path
            except Exception as e:
                results['questionnaire'] = None
                print(f"Error generating questionnaire: {e}")
        
        # 実験台本（手順がある場合）
        if form_data.get('procedures'):
            try:
                path = self.generate_experiment_script(
                    title=form_data.get('research_title', ''),
                    procedures=form_data['procedures'],
                    form_data=form_data
                )
                results['experiment_script'] = path
            except Exception as e:
                results['experiment_script'] = None
                print(f"Error generating experiment script: {e}")
        
        return results
    
    def create_zip(self, generated_files: Dict[str, Optional[Path]]) -> Path:
        """生成されたファイルをZIPにまとめる"""
        zip_path = self.output_dir / "倫理審査書類一式.zip"
        
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            for doc_type, file_path in generated_files.items():
                if file_path and file_path.exists():
                    zf.write(file_path, file_path.name)
        
        return zip_path
    
    def get_zip_bytes(self, generated_files: Dict[str, Optional[Path]]) -> bytes:
        """生成されたファイルをZIPのバイト列として返す"""
        buffer = io.BytesIO()
        
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            for doc_type, file_path in generated_files.items():
                if file_path and file_path.exists():
                    zf.write(file_path, file_path.name)
        
        buffer.seek(0)
        return buffer.getvalue()

