from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from app.services.official_docx_renderer import first_nonempty, get_path, get_template_path


class XlsxTemplateValidationError(RuntimeError):
    pass


OLD_PARTICIPANT_DENYLIST = [
    "山田　貴義",
    "青栁匠悟",
    "川村　誉羅",
    "山内由大",
    "倉友乃康",
    "古茂田隆之",
    "石島　諒一",
    "原馬 誓一郎",
    "藤田陶子",
    "王　兆龍",
    "s2330183@u.tsukuba.ac.jp",
    "s2220697@u.tsukuba.ac.jp",
    "s2320503@u.tsukuba.ac.jp",
]


def _format_amount(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    if isinstance(value, (int, float)):
        return f"{int(value):,}"
    text = str(value).strip()
    if text.replace(",", "").isdigit():
        return f"{int(text.replace(',', '')):,}"
    return text


def _workbook_text(workbook: Workbook) -> str:
    values: list[str] = []
    for sheet in workbook.worksheets:
        values.append(sheet.title)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    values.append(str(cell.value))
    return "\n".join(values)


def _validate_participant_list(workbook: Workbook) -> None:
    text = _workbook_text(workbook)
    leftovers = [value for value in OLD_PARTICIPANT_DENYLIST if value in text]
    if "Sheet1" in workbook.sheetnames:
        leftovers.append("Sheet1 with prefilled participant data")
    if leftovers:
        raise XlsxTemplateValidationError(f"Rendered participant list contains stale values: {leftovers}")


def render_participant_list(context: dict[str, Any], output_path: Path) -> Path:
    workbook = load_workbook(get_template_path("participant_list"))
    official_sheet = workbook["公式"] if "公式" in workbook.sheetnames else workbook.worksheets[0]

    for sheet in list(workbook.worksheets):
        if sheet is not official_sheet:
            workbook.remove(sheet)
    official_sheet.title = "公式"
    workbook.active = 0

    reward_type = first_nonempty(get_path(context, "reward.type"), "Amazonギフト券（メールタイプ）")
    reward_amount = _format_amount(get_path(context, "reward.amount"))

    for row in range(4, 36):
        for column in ["B", "C", "D", "G", "J"]:
            official_sheet[f"{column}{row}"] = None
        official_sheet[f"K{row}"] = reward_type
        official_sheet[f"M{row}"] = reward_amount

    official_sheet["I2"] = '=COUNTIF(B4:B35, "<>")'

    _validate_participant_list(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def render_official_xlsx_template(template_key: str, context: dict[str, Any], output_path: Path) -> Path:
    if template_key == "participant_list":
        return render_participant_list(context, output_path)
    raise ValueError(f"Unsupported official xlsx template: {template_key}")
