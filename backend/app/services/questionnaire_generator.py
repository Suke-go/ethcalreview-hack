from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pydantic import BaseModel, Field

from app.logger import get_logger
from app.services.llm_client import LLMClient


logger = get_logger(__name__)

ResponseType = Literal["single_choice", "multiple_choice", "likert", "free_text", "numeric"]
TimingType = Literal["pre", "per_stimulus", "post_condition", "post_experiment"]


class ScaleSpec(BaseModel):
    min: int = 1
    max: int = 7
    min_label: str = "全くそう思わない"
    max_label: str = "非常にそう思う"


class SurveyItem(BaseModel):
    item_id: str
    question: str
    construct: str
    response_type: ResponseType
    options: list[str] = Field(default_factory=list)
    scale: ScaleSpec | None = None
    required: bool = True


class QuestionnaireBlock(BaseModel):
    block_id: str
    title: str
    timing: TimingType
    items: list[SurveyItem] = Field(default_factory=list)


class QuestionnaireSpec(BaseModel):
    title: str
    description: str
    blocks: list[QuestionnaireBlock] = Field(default_factory=list)


QUESTIONNAIRE_SYSTEM_INSTRUCTION = """
あなたは研究倫理申請に添付するアンケート用紙を設計する専門家です。
あなたの仕事は、研究目的・実験条件・評価指標に対応した質問項目をJSONで作ることです。

制約:
- docxのレイアウトは固定処理が担当するため、本文と質問項目だけを出力してください。
- ユーザー入力をそのまま長文コピーしないでください。
- 刺激内容が未確定の理解確認問題は、刺激ごとに差し替えられるテンプレート質問として作ってください。
- NASA-TLXなど既存尺度の完全転載は避け、倫理申請用の簡易主観評価項目として作ってください。
- 個人を過度に識別する質問は避けてください。
- 用語の制約：質問文や説明文で研究に協力する人を指す場合は必ず「研究対象者」または「参加者」と表記してください。それ以外の旧来の呼称（健康状態を含意する語や、実験の語を含む対象側の旧来の呼称など）は使用しないでください。
- 出力はJSONのみです。
"""


def _as_text(value: Any) -> str:
    if isinstance(value, list):
        return "、".join(str(item) for item in value if str(item).strip())
    return str(value or "")


def _infer_conditions(form_data: dict[str, Any]) -> list[str]:
    # 明示的に与えられた実験条件のみを使う。特定研究（字幕等）の条件を勝手に注入しない。
    # 入力に条件が無ければ空のままとし、条件に依存しない汎用の主観評価項目を作る。
    explicit = form_data.get("conditions") or form_data.get("condition_names")
    if isinstance(explicit, list) and explicit:
        return [str(item) for item in explicit if str(item).strip()]
    if isinstance(explicit, str) and explicit.strip():
        return [part.strip() for part in re.split(r"[、,\n]", explicit) if part.strip()]
    return []


def _context_for_prompt(form_data: dict[str, Any], questionnaire_type: str) -> dict[str, Any]:
    return {
        "questionnaire_type": questionnaire_type,
        "research_title": form_data.get("title") or form_data.get("research_title") or "",
        "purpose": form_data.get("purpose") or form_data.get("research_purpose") or "",
        "method": form_data.get("methodology") or form_data.get("research_method") or "",
        "target_participants": form_data.get("targetDescription") or form_data.get("target_participants") or "",
        "duration_minutes": form_data.get("duration") or form_data.get("duration_minutes") or "",
        "conditions": _infer_conditions(form_data),
        "devices": form_data.get("devices", []),
        "risks": form_data.get("risks", []),
        "data_types": form_data.get("dataTypes") or form_data.get("data_types") or [],
    }


def _questionnaire_prompt(form_data: dict[str, Any], questionnaire_type: Literal["pre", "post"]) -> str:
    context = _context_for_prompt(form_data, questionnaire_type)
    if questionnaire_type == "pre":
        intent = """
事前アンケートを作成してください。特定の研究テーマを想定せず、研究計画コンテキストに基づいて項目を作ってください。
含めるべき観点:
- 参加条件の確認（対象者条件・除外基準に対応）
- 年齢区分など最小限の属性（過度に個人を特定しない範囲）
- 本研究の課題遂行に関係する経験・背景（研究計画に書かれている範囲のみ）
- 課題遂行やリスクに関係する体調・感覚などの自己申告
"""
    else:
        intent = """
実験中または実験後に使うアンケートを作成してください。特定の研究テーマを想定せず、研究計画コンテキスト（目的・方法・conditions・data_types・評価指標）に基づいて項目を具体化してください。入力に無い構成は作らないでください。
含めるべき観点:
- 各条件・各刺激に対する主観評価（研究計画の評価指標・従属変数に対応）
- 課題の理解度・難易度など、研究計画上測定する内容の確認（該当する場合のみ）
- 簡易的な主観的負担（必要に応じて）
- 実験全体の比較評価と自由記述
"""

    return f"""
以下の研究計画に基づいて、研究倫理申請に添付できるアンケート用紙案を生成してください。

{intent}

研究計画コンテキスト:
{json.dumps(context, ensure_ascii=False, indent=2)}

出力JSON:
{{
  "questionnaire": {{
    "title": "string",
    "description": "string",
    "blocks": [
      {{
        "block_id": "string",
        "title": "string",
        "timing": "pre | per_stimulus | post_condition | post_experiment",
        "items": [
          {{
            "item_id": "string",
            "question": "string",
            "construct": "この質問が測るもの",
            "response_type": "single_choice | multiple_choice | likert | free_text | numeric",
            "options": ["string"],
            "scale": {{
              "min": 1,
              "max": 7,
              "min_label": "string",
              "max_label": "string"
            }},
            "required": true
          }}
        ]
      }}
    ]
  }}
}}
"""


def _fallback_pre_questionnaire() -> QuestionnaireSpec:
    # 特定研究を想定しない汎用の事前アンケート（LLM生成に失敗した場合の最小構成）
    return QuestionnaireSpec(
        title="事前アンケート",
        description="研究参加前に、参加条件と課題遂行に関係する状態を確認します。",
        blocks=[
            QuestionnaireBlock(
                block_id="pre_background",
                title="参加前確認",
                timing="pre",
                items=[
                    SurveyItem(
                        item_id="pre_age",
                        question="年齢区分を選択してください。",
                        construct="参加条件",
                        response_type="single_choice",
                        options=["18-19歳", "20-29歳", "30-39歳", "40歳以上", "回答しない"],
                    ),
                    SurveyItem(
                        item_id="pre_condition",
                        question="本研究の課題を行う上で支障となる健康上・身体上の事情はありますか。",
                        construct="参加条件の自己申告",
                        response_type="single_choice",
                        options=["特にない", "少しある", "ある", "回答しない"],
                    ),
                    SurveyItem(
                        item_id="pre_physical",
                        question="本日の体調はいかがですか。",
                        construct="当日の体調",
                        response_type="likert",
                        scale=ScaleSpec(min_label="非常に悪い", max_label="非常に良い"),
                    ),
                ],
            )
        ],
    )


def _fallback_post_questionnaire(conditions: list[str]) -> QuestionnaireSpec:
    condition_text = "、".join(conditions) if conditions else "各条件"
    # 特定研究を想定しない汎用の事後アンケート。条件が与えられた場合のみ条件比較を加える。
    per_task_items = [
        SurveyItem(
            item_id="task_difficulty",
            question="課題はどの程度難しく感じましたか。",
            construct="主観的難易度",
            response_type="likert",
            scale=ScaleSpec(min_label="全く難しくなかった", max_label="非常に難しかった"),
        ),
        SurveyItem(
            item_id="task_load",
            question="課題を行う際に、頭を使う負担を感じましたか。",
            construct="主観的負担",
            response_type="likert",
            scale=ScaleSpec(min_label="全く負担を感じなかった", max_label="非常に負担を感じた"),
        ),
    ]
    overall_items: list[SurveyItem] = []
    if conditions:
        overall_items.append(
            SurveyItem(
                item_id="post_best_condition",
                question="最も良かったと感じた条件を選択してください。",
                construct="条件比較",
                response_type="single_choice",
                options=conditions,
            )
        )
    overall_items.append(
        SurveyItem(
            item_id="post_free",
            question="実験全体について、気づいた点があれば自由に記入してください。",
            construct="自由記述",
            response_type="free_text",
        )
    )
    description = (
        f"{condition_text}での課題遂行時の主観評価と全体的な所感を確認します。"
        if conditions
        else "課題遂行時の主観評価と全体的な所感を確認します。"
    )
    return QuestionnaireSpec(
        title="実験後アンケート",
        description=description,
        blocks=[
            QuestionnaireBlock(
                block_id="per_task",
                title="課題後の質問",
                timing="post_condition",
                items=per_task_items,
            ),
            QuestionnaireBlock(
                block_id="post_overall",
                title="実験全体について",
                timing="post_experiment",
                items=overall_items,
            ),
        ],
    )


class QuestionnaireGenerator:
    def __init__(self, llm_client: LLMClient, lab_defaults: dict[str, Any]):
        self.llm = llm_client
        self.defaults = lab_defaults

    async def generate_pre_questionnaire(self, form_data: dict[str, Any], output_dir: Path) -> Path:
        spec = await self._generate_spec(form_data, "pre")
        return self._save_outputs(spec, output_dir, "事前アンケート")

    async def generate_post_questionnaire(self, form_data: dict[str, Any], output_dir: Path) -> Path:
        spec = await self._generate_spec(form_data, "post")
        return self._save_outputs(spec, output_dir, "実験後アンケート")

    async def _generate_spec(
        self,
        form_data: dict[str, Any],
        questionnaire_type: Literal["pre", "post"],
    ) -> QuestionnaireSpec:
        try:
            generated = await self.llm.generate_json(
                _questionnaire_prompt(form_data, questionnaire_type),
                QUESTIONNAIRE_SYSTEM_INSTRUCTION,
            )
            payload = generated.get("questionnaire", generated)
            spec = QuestionnaireSpec.model_validate(payload)
            if spec.blocks and any(block.items for block in spec.blocks):
                return spec
        except Exception as exc:
            logger.warning(f"アンケートJSON生成に失敗しました。fallbackを使用します: {exc}")

        if questionnaire_type == "pre":
            return _fallback_pre_questionnaire()
        return _fallback_post_questionnaire(_infer_conditions(form_data))

    def _save_outputs(self, spec: QuestionnaireSpec, output_dir: Path, basename: str) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        json_path = output_dir / f"{basename}.json"
        json_path.write_text(spec.model_dump_json(indent=2), encoding="utf-8")

        docx_path = output_dir / f"{basename}.docx"
        self._generate_docx(spec, docx_path)
        logger.info(f"{basename}生成完了: {docx_path.name}, {json_path.name}")
        return docx_path

    def _generate_docx(self, spec: QuestionnaireSpec, output_path: Path) -> None:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Yu Gothic"
        style.font.size = Pt(10.5)

        title_para = doc.add_paragraph()
        run = title_para.add_run(spec.title)
        run.bold = True
        run.font.size = Pt(14)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if spec.description:
            doc.add_paragraph(spec.description)
        doc.add_paragraph("回答に正解・不正解はありません。普段どおりに回答してください。")

        question_number = 1
        for block in spec.blocks:
            heading = doc.add_paragraph()
            heading_run = heading.add_run(block.title)
            heading_run.bold = True
            heading_run.font.size = Pt(12)
            if block.timing == "per_stimulus":
                doc.add_paragraph("このブロックは、各動画または各条件の後に繰り返して使用します。")

            for item in block.items:
                paragraph = doc.add_paragraph()
                paragraph.add_run(f"Q{question_number}. {item.question}").bold = True
                if item.construct:
                    doc.add_paragraph(f"測定内容: {item.construct}")
                self._render_response_area(doc, item)
                doc.add_paragraph()
                question_number += 1

        doc.save(output_path)

    def _render_response_area(self, doc: Document, item: SurveyItem) -> None:
        if item.response_type == "likert":
            scale = item.scale or ScaleSpec()
            self._add_likert_scale(doc, scale)
        elif item.response_type in {"single_choice", "multiple_choice"}:
            for option in item.options:
                doc.add_paragraph(f"　□ {option}")
        elif item.response_type == "numeric":
            doc.add_paragraph("　回答: _______________")
        else:
            doc.add_paragraph("　" + "_" * 50)
            doc.add_paragraph("　" + "_" * 50)

    def _add_likert_scale(self, doc: Document, scale: ScaleSpec) -> None:
        min_value = int(scale.min)
        max_value = int(scale.max)
        values = list(range(min_value, max_value + 1))
        table = doc.add_table(rows=2, cols=len(values))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for index, value in enumerate(values):
            label = ""
            if index == 0:
                label = scale.min_label
            elif index == len(values) - 1:
                label = scale.max_label
            table.rows[0].cells[index].text = label
            table.rows[0].cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            table.rows[1].cells[index].text = f"□ {value}"
            table.rows[1].cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


async def generate_questionnaire(
    form_data: dict[str, Any],
    output_dir: Path,
    llm_client: LLMClient,
    lab_defaults: dict[str, Any],
    questionnaire_type: Literal["pre", "post", "both"] = "both",
) -> list[Path]:
    generator = QuestionnaireGenerator(llm_client, lab_defaults)
    paths: list[Path] = []

    if questionnaire_type in {"pre", "both"}:
        paths.append(await generator.generate_pre_questionnaire(form_data, output_dir))
    if questionnaire_type in {"post", "both"}:
        paths.append(await generator.generate_post_questionnaire(form_data, output_dir))

    return paths
