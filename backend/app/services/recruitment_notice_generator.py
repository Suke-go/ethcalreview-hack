"""
募集案内文 DOCX生成

研究参加者募集の案内文をdocx形式で生成します。
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
from typing import Dict, Any, List
from datetime import datetime


def generate_recruitment_notice(
    form_data: Dict[str, Any],
    output_dir: Path,
    lab_defaults: Dict[str, Any]
) -> Path:
    """
    募集案内文を生成
    
    Args:
        form_data: フォームデータ
        output_dir: 出力ディレクトリ
        lab_defaults: 研究室デフォルト設定
    
    Returns:
        生成されたファイルのパス
    """
    doc = Document()
    
    # スタイル設定（11pt統一）
    style = doc.styles['Normal']
    style.font.name = 'Yu Gothic'
    style.font.size = Pt(11)
    
    # ========================================
    # タイトル
    # ========================================
    title = doc.add_paragraph()
    title_run = title.add_run(f"{form_data.get('research_title', '')}における")
    title_run.font.size = Pt(11)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run('実験参加者募集')
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(14)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # ========================================
    # 研究概要
    # ========================================
    overview = form_data.get('research_overview', '')
    if overview:
        doc.add_paragraph(overview)
        doc.add_paragraph()
    
    # 区切り線
    doc.add_paragraph('─' * 40)
    
    # ========================================
    # 参加条件
    # ========================================
    conditions_heading = doc.add_paragraph()
    conditions_heading.add_run('[参加条件]').bold = True
    
    participant_criteria = form_data.get('participant_criteria', [])
    if isinstance(participant_criteria, list):
        for criteria in participant_criteria:
            doc.add_paragraph(criteria, style='List Bullet')
    else:
        doc.add_paragraph(str(participant_criteria))
    
    # 除外基準
    exclusion = form_data.get('exclusion_criteria', [])
    if exclusion:
        doc.add_paragraph('以下のいずれにも該当しない方')
        if isinstance(exclusion, list):
            for item in exclusion:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(item).font.size = Pt(11)
        else:
            doc.add_paragraph(str(exclusion))
    
    doc.add_paragraph()
    
    # ========================================
    # 実験の内容
    # ========================================
    content_heading = doc.add_paragraph()
    content_heading.add_run('[実験の内容]').bold = True
    
    intro = form_data.get('experiment_intro', '実験では、以下の手順で評価していただきます。')
    doc.add_paragraph(intro)
    
    steps = form_data.get('experiment_steps', [])
    if isinstance(steps, list):
        for i, step in enumerate(steps, 1):
            step_text = step if isinstance(step, str) else step.get('description', '')
            p = doc.add_paragraph()
            p.add_run(f"{i}. {step_text}").font.size = Pt(11)
    
    doc.add_paragraph()
    
    # ========================================
    # 謝礼
    # ========================================
    reward_heading = doc.add_paragraph()
    reward_heading.add_run('[謝礼]').bold = True
    
    reward_amount = form_data.get('reward_amount', '')
    reward_type = form_data.get('reward_type', 'Amazonギフトカード（Eメールタイプ）')
    
    if reward_amount:
        doc.add_paragraph(
            f"実験終了時に、謝金として{reward_type}{reward_amount}円分をお支払いいたします。"
        )
    
    doc.add_paragraph()
    
    # ========================================
    # ご注意
    # ========================================
    caution_heading = doc.add_paragraph()
    caution_heading.add_run('[ご注意]').bold = True
    
    cautions = form_data.get('cautions', '')
    if cautions:
        doc.add_paragraph(cautions)
    else:
        # デフォルトの注意事項
        duration = form_data.get('duration_minutes', 45)
        doc.add_paragraph(
            f"本実験を実施する前に十分な説明を行います。"
            f"実験中に不快感を感じた場合は、いつでも実験を中断・中止することが可能です。"
        )
        doc.add_paragraph(f"説明と準備を含め、所要時間は約{duration}分程度です。")
    
    doc.add_paragraph()
    doc.add_paragraph("ご興味のある方は、スタッフにお声がけください。")
    
    # 保存
    output_path = output_dir / "募集案内文.docx"
    doc.save(output_path)
    
    return output_path
