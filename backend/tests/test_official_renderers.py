from __future__ import annotations

from pathlib import Path
from shutil import rmtree

import pytest
from docx import Document
from openpyxl import load_workbook

from app.config import app_config, load_user_settings
from app.services.form_context_builder import build_generation_context
from app.services.generation_validator import error_issues, validate_generation_context
from app.services.official_docx_renderer import all_paragraphs, render_official_template
from app.services.official_xlsx_renderer import render_official_xlsx_template
from app.services.preset_manager import load_preset_bundle


@pytest.fixture
def output_dir() -> Path:
    path = Path(".pytest_tmp") / "official_renderers"
    if path.exists():
        rmtree(path)
    path.mkdir(parents=True)
    try:
        yield path
    finally:
        if path.exists():
            rmtree(path)


@pytest.fixture
def generation_context() -> dict:
    settings = load_user_settings(app_config)
    presets = load_preset_bundle(app_config)
    form_data = {
        "title": "Official Render Test",
        "purpose": "Purpose",
        "methodology": "Task and questionnaire",
        "targetDescription": "Adult participants",
        "inclusionCriteria": ["Adults"],
        "exclusionCriteria": ["None"],
        "expectedParticipants": 10,
        "participantsJustification": "Needed for validation",
        "procedures": ["Task and questionnaire"],
        "risks": ["Minor fatigue"],
        "riskCountermeasures": ["Rest"],
        "duration": 60,
        "rewardAmount": 1000,
        "app_config": {
            "applicationType": "new",
            "researchPeriodEndText": "2027年3月31日",
            "domainName": "システム情報系",
            "domainHeadName": "矢野 博明",
            "subInvestigators": [
                {
                    "affiliation": "University",
                    "position": "Assistant Professor",
                    "name": "Taro Yamada",
                    "tel": "029-000-0000",
                }
            ],
            "facilityName": "3M211",
            "fundingSource": "Operating Fund",
            "fundingPI": "Keiichi Zempo",
            "fundingProjectName": "Validation Study",
            "dataTypes": "Reaction time, questionnaire",
            "storageLocation": "Lab 3M211",
            "dataManager": "Keiichi Zempo",
            "managementMethod": "Encrypted storage",
            "disposalMethod": "Secure deletion",
            "videoRecording": False,
            "invasiveness": False,
            "conflictOfInterest": False,
        },
    }
    context = build_generation_context(form_data, settings, presets)
    errors = error_issues(validate_generation_context(context))
    assert errors == []
    return context


def document_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in all_paragraphs(document))


@pytest.mark.parametrize(
    "template_key",
    [
        "application_form",
        "consent_form",
        "consent_withdrawal",
        "honorarium_rationale",
    ],
)
def test_official_docx_templates_render_required_values(
    template_key: str,
    generation_context: dict,
    output_dir: Path,
) -> None:
    path = render_official_template(template_key, generation_context, output_dir / f"{template_key}.docx")
    text = document_text(path)

    assert "Official Render Test" in text
    if template_key == "honorarium_rationale":
        assert "Operating Fund" in text
        assert "1,000" in text
        assert "10,000" in text
    if template_key in {"consent_form", "consent_withdrawal"}:
        assert "Taro Yamada" in text
        assert "029-000-0000" in text


def test_consent_form_fills_back_side_overview(generation_context: dict, output_dir: Path) -> None:
    """同意書裏面（別紙「研究の概要について」①②③）が context から流し込まれること。"""
    path = render_official_template("consent_form", generation_context, output_dir / "consent_form.docx")
    text = document_text(path)

    # ① 研究の概要：目的・方法・参加者条件・所要時間・実験手順
    assert "Purpose" in text  # [目的]
    assert "Task and questionnaire" in text  # [方法] / 実験手順
    assert "Adult participants" in text  # [参加者条件]
    assert "【実験手順】" in text  # procedures が裏面に載る
    assert "約60分" in text  # [所要時間]
    # ③ 個人情報保護：任意性の固定文と撤回期限
    assert "研究への参加は任意であり" in text
    assert "同意書署名の日から90日後" in text
    # テンプレに残っていた旧サンプル本文（前任者の管理方法説明）が消えていること
    assert "特定のPCのみを用います" not in text


def test_application_form_uses_260422_layout(generation_context: dict, output_dir: Path) -> None:
    """01-1 申請書が 260422 新公式（1.x/2.x/3.x…）として描画されること。"""
    path = render_official_template("application_form", generation_context, output_dir / "application_form.docx")
    text = document_text(path)

    # 新章番号の見出しが存在し、旧 1〜14 様式に作り替えられていないこと
    assert "１研究計画の概要" in text
    assert "2　取得データに関する情報" in text
    assert "3. 実験対象者" in text
    # チェックボックスが排他的に設定されること（新規申請・単独施設・侵襲なし）
    assert "■新規申請" in text
    assert "■a. 筑波大学単独施設での研究" in text
    assert "侵襲性（□軽微でない　□軽微　■無）" in text
    # 値が所定位置に入ること
    assert "1.1　課題名　Official Render Test" in text
    assert "3.1 期間内に実施する対象者の人数の見積もり：10人" in text
    assert "Taro Yamada" in text  # 1.4 実施分担者として挿入
    # 旧様式の番号体系が残っていないこと
    assert "１４　実施責任者の問い合わせ先" not in text


def test_participant_list_keeps_only_official_sheet_and_removes_prefilled_personal_data(
    generation_context: dict,
    output_dir: Path,
) -> None:
    path = render_official_xlsx_template("participant_list", generation_context, output_dir / "participant_list.xlsx")
    workbook = load_workbook(path, data_only=False)

    assert workbook.sheetnames == ["公式"]
    sheet = workbook["公式"]
    assert sheet["K4"].value
    assert str(sheet["M4"].value) in {"1000", "1,000"}

    workbook_text = "\n".join(
        str(cell.value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "s2330183@u.tsukuba.ac.jp" not in workbook_text
    assert "Sheet1" not in workbook.sheetnames


def test_reward_total_amount_is_derived_for_official_templates(generation_context: dict) -> None:
    assert generation_context["reward"]["amount"] == 1000
    assert generation_context["participants"]["count"] == 10
    assert generation_context["reward"]["total_amount"] == 10000
    assert generation_context["participant_list"]["total_reward"] == 10000
