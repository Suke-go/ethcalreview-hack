"""Tests for participant_list xlsx rendering (P2-7).

Verifies that:
(a) The output workbook contains only the '公式' sheet.
(b) Sample personal data (e.g. s2330183@u.tsukuba.ac.jp) is absent from output.
(c) The reward amount is written to M4 (numeric or comma-formatted string).
(d) The reward type is written to K4.
(e) New column headers added in the refer-sync (実験日, 実施時間) are present in row 3.
"""
from __future__ import annotations

from pathlib import Path
from shutil import rmtree

import pytest
from openpyxl import load_workbook

from app.config import app_config, load_user_settings
from app.services.form_context_builder import build_generation_context
from app.services.generation_validator import error_issues, validate_generation_context
from app.services.official_xlsx_renderer import render_official_xlsx_template
from app.services.preset_manager import load_preset_bundle


@pytest.fixture
def output_dir() -> Path:
    path = Path(".pytest_tmp") / "xlsx_render"
    if path.exists():
        rmtree(path)
    path.mkdir(parents=True)
    try:
        yield path
    finally:
        if path.exists():
            rmtree(path)


@pytest.fixture
def generation_context() -> dict:
    settings = load_user_settings(app_config)
    presets = load_preset_bundle(app_config)
    form_data = {
        "title": "XlsxRender Test",
        "purpose": "Purpose",
        "methodology": "Task and questionnaire",
        "targetDescription": "Adult participants",
        "inclusionCriteria": ["Adults"],
        "exclusionCriteria": ["None"],
        "expectedParticipants": 10,
        "participantsJustification": "Needed for validation",
        "procedures": ["Task and questionnaire"],
        "risks": ["Minor fatigue"],
        "riskCountermeasures": ["Rest"],
        "duration": 60,
        "rewardAmount": 1000,
        "app_config": {
            "applicationType": "new",
            "researchPeriodEndText": "2027年3月31日",
            "domainName": "システム情報系",
            "domainHeadName": "矢野 博明",
            "subInvestigators": [
                {
                    "affiliation": "University",
                    "position": "Assistant Professor",
                    "name": "Taro Yamada",
                    "tel": "029-000-0000",
                }
            ],
            "facilityName": "3M211",
            "fundingSource": "Operating Fund",
            "fundingPI": "Keiichi Zempo",
            "fundingProjectName": "Validation Study",
            "dataTypes": "Reaction time, questionnaire",
            "storageLocation": "Lab 3M211",
            "dataManager": "Keiichi Zempo",
            "managementMethod": "Encrypted storage",
            "disposalMethod": "Secure deletion",
            "videoRecording": False,
            "invasiveness": False,
            "conflictOfInterest": False,
        },
    }
    context = build_generation_context(form_data, settings, presets)
    errors = error_issues(validate_generation_context(context))
    assert errors == []
    return context


def _workbook_all_text(path: Path) -> str:
    wb = load_workbook(path, data_only=False)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(ws.title)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    parts.append(str(cell.value))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# (a) Only 公式 sheet is present
# ---------------------------------------------------------------------------
def test_participant_list_has_only_official_sheet(generation_context: dict, output_dir: Path) -> None:
    path = render_official_xlsx_template(
        "participant_list", generation_context, output_dir / "participant_list.xlsx"
    )
    wb = load_workbook(path)
    assert wb.sheetnames == ["公式"], f"Unexpected sheets: {wb.sheetnames}"


# ---------------------------------------------------------------------------
# (b) No sample personal data in output
# ---------------------------------------------------------------------------
SAMPLE_EMAILS = [
    "s2330183@u.tsukuba.ac.jp",
    "s2220697@u.tsukuba.ac.jp",
    "s2320503@u.tsukuba.ac.jp",
]
SAMPLE_NAMES = [
    "山田　貴義",
    "青栁匠悟",
    "川村　誉羅",
]


@pytest.mark.parametrize("sample_value", SAMPLE_EMAILS + SAMPLE_NAMES)
def test_participant_list_no_sample_personal_data(
    sample_value: str, generation_context: dict, output_dir: Path
) -> None:
    path = render_official_xlsx_template(
        "participant_list",
        generation_context,
        output_dir / f"participant_list_{hash(sample_value)}.xlsx",
    )
    text = _workbook_all_text(path)
    assert sample_value not in text, f"Sample value '{sample_value}' found in output"


# ---------------------------------------------------------------------------
# (c) Reward amount is written to M4
# ---------------------------------------------------------------------------
def test_participant_list_reward_amount_in_M4(generation_context: dict, output_dir: Path) -> None:
    path = render_official_xlsx_template(
        "participant_list", generation_context, output_dir / "participant_list.xlsx"
    )
    wb = load_workbook(path, data_only=False)
    sheet = wb["公式"]
    m4_value = str(sheet["M4"].value)
    assert m4_value in {"1000", "1,000"}, f"M4 value was: {m4_value!r}"


# ---------------------------------------------------------------------------
# (d) Reward type is written to K4 (truthy)
# ---------------------------------------------------------------------------
def test_participant_list_reward_type_in_K4(generation_context: dict, output_dir: Path) -> None:
    path = render_official_xlsx_template(
        "participant_list", generation_context, output_dir / "participant_list.xlsx"
    )
    wb = load_workbook(path, data_only=False)
    sheet = wb["公式"]
    assert sheet["K4"].value, f"K4 should contain reward type, got: {sheet['K4'].value!r}"


# ---------------------------------------------------------------------------
# (e) Column headers match the refer-synced layout (実験日 at E3, 実施時間 at F3)
# ---------------------------------------------------------------------------
def test_participant_list_header_columns_synced_to_refer(
    generation_context: dict, output_dir: Path
) -> None:
    path = render_official_xlsx_template(
        "participant_list", generation_context, output_dir / "participant_list.xlsx"
    )
    wb = load_workbook(path, data_only=False)
    sheet = wb["公式"]

    # Core columns (always required)
    assert sheet["A3"].value == "No."
    assert sheet["B3"].value == "氏名"
    assert sheet["C3"].value == "所属"
    assert sheet["D3"].value == "メールアドレス"
    assert sheet["K3"].value == "謝礼形式"
    assert sheet["M3"].value == "金額"

    # Columns synced from refer template
    assert sheet["E3"].value == "実験日", f"E3 should be '実験日', got: {sheet['E3'].value!r}"
    assert sheet["F3"].value == "実施時間", f"F3 should be '実施時間', got: {sheet['F3'].value!r}"
