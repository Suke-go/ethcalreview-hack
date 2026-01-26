#!/usr/bin/env python3
"""
テンプレート自動置換スクリプト

サンプルDOCXファイルをJinja2テンプレートに変換します。
特定の値をプレースホルダー（{{ variable }}）に置換します。
"""

from docx import Document
from pathlib import Path
import sys
from typing import Dict, List


def replace_text_in_paragraph(paragraph, replacements: Dict[str, str]) -> bool:
    """
    パラグラフ内のテキストを置換（フォーマット保持）
    
    Args:
        paragraph: パラグラフオブジェクト
        replacements: 置換マッピング {元テキスト: 新テキスト}
    
    Returns:
        置換が行われた場合True
    """
    replaced = False
    full_text = paragraph.text
    
    for old_text, new_text in replacements.items():
        if old_text in full_text:
            # Run単位で置換してフォーマットを保持
            inline = paragraph.runs
            for run in inline:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)
                    replaced = True
    
    return replaced


def replace_text_in_table(table, replacements: Dict[str, str]) -> int:
    """
    表内のテキストを置換
    
    Args:
        table: 表オブジェクト
        replacements: 置換マッピング
    
    Returns:
        置換回数
    """
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if replace_text_in_paragraph(paragraph, replacements):
                    count += 1
    return count


def templatize_document(
    input_path: Path,
    output_path: Path,
    replacements: Dict[str, str]
) -> Dict[str, int]:
    """
    DOCXファイルをテンプレート化
    
    Args:
        input_path: 入力ファイルパス
        output_path: 出力ファイルパス
        replacements: 置換マッピング
    
    Returns:
        統計情報 {"paragraphs": 置換パラグラフ数, "tables": 置換セル数}
    """
    print(f"[INPUT]  {input_path}")
    print(f"[OUTPUT] {output_path}")
    print(f"[REPLACE] {len(replacements)} patterns")
    print("-" * 60)
    
    # ドキュメント読み込み
    doc = Document(input_path)
    
    stats = {"paragraphs": 0, "tables": 0}
    
    # パラグラフ内のテキスト置換
    for para in doc.paragraphs:
        if replace_text_in_paragraph(para, replacements):
            stats["paragraphs"] += 1
    
    # 表内のテキスト置換
    for table in doc.tables:
        stats["tables"] += replace_text_in_table(table, replacements)
    
    # 保存
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    
    print(f"[OK] Replaced {stats['paragraphs']} paragraphs, {stats['tables']} table cells")
    return stats


# ========================================
# 置換マッピング定義
# ========================================

# 同意書の置換マッピング
CONSENT_FORM_REPLACEMENTS = {
    # 研究代表者
    "善甫 啓一": "{{ pi_name }}",
    "筑波大学 システム情報系": "{{ pi_affiliation }}",
    "029-853-5338": "{{ pi_tel }}",
    
    # 実施担当者
    "清水 啓斗": "{{ conductor_name }}",
    "029-853-6185": "{{ conductor_tel }}",
    
    # 倫理委員会
    "筑波大学 システム情報系 研究倫理委員会事務局": "{{ ethics_committee }}",
    "システム情報エリア支援室": "{{ ethics_office }}",
    "029-853-4989": "{{ ethics_tel }}",
    
    # 研究内容（これらはサンプル文を一括置換）
    # 実際の研究目的・方法は各自のサンプルファイルに合わせて調整
}

# 同意撤回書の置換マッピング
CONSENT_WITHDRAWAL_REPLACEMENTS = {
    "善甫 啓一": "{{ pi_name }}",
    "筑波大学 システム情報系": "{{ pi_affiliation }}",
    "029-853-5338": "{{ pi_tel }}",
    "清水 啓斗": "{{ conductor_name }}",
    "029-853-6185": "{{ conductor_tel }}",
    "筑波大学 システム情報系 研究倫理委員会事務局": "{{ ethics_committee }}",
    "029-853-4989": "{{ ethics_tel }}",
}


def main():
    """メイン処理"""
    project_root = Path(__file__).parent.parent
    
    # ========================================
    # 1. 同意書のテンプレート化
    # ========================================
    print("\n" + "=" * 60)
    print("[1] 同意書のテンプレート化")
    print("=" * 60)
    
    consent_input = project_root / "sample" / "03_同意書（shimizu）_v3.docx"
    consent_output = project_root / "backend" / "templates" / "03_同意書_template.docx"
    
    if consent_input.exists():
        templatize_document(
            consent_input,
            consent_output,
            CONSENT_FORM_REPLACEMENTS
        )
    else:
        print(f"[ERROR] File not found: {consent_input}")
    
    # ========================================
    # 2. 同意撤回書のテンプレート化
    # ========================================
    print("\n" + "=" * 60)
    print("[2] 同意撤回書のテンプレート化")
    print("=" * 60)
    
    withdrawal_input = project_root / "sample" / "04_同意撤回書.ver1 1.docx"
    withdrawal_output = project_root / "backend" / "templates" / "04_同意撤回書_template.docx"
    
    if withdrawal_input.exists():
        templatize_document(
            withdrawal_input,
            withdrawal_output,
            CONSENT_WITHDRAWAL_REPLACEMENTS
        )
    else:
        print(f"[ERROR] File not found: {withdrawal_input}")
    
    # ========================================
    # 完了メッセージ
    # ========================================
    print("\n" + "=" * 60)
    print("[DONE] Template generation completed")
    print("=" * 60)
    print(f"\nGenerated templates:")
    print(f"  - {consent_output}")
    print(f"  - {withdrawal_output}")
    print(f"\nNext steps:")
    print(f"  1. Open templates in Word and verify")
    print(f"  2. Check placeholders with DocxTemplate")


if __name__ == "__main__":
    main()
